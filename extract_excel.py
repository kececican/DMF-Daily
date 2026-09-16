#!/usr/bin/env python3
"""DMF Daily — Excel → dashboard/data.json ayıklayıcı.

Excel yolu DMF_EXCEL ortam değişkeni ile verilebilir:
    DMF_EXCEL=/yol/dosya.xlsx python3 extract_excel.py
"""

import os
import json
from datetime import datetime
from collections import defaultdict

BASE_DIR   = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.environ.get(
    'DMF_EXCEL',
    os.path.join(BASE_DIR, 'data', 'DMF_Daily_Followup_New_Version_2026.xlsx'),
)
CACHE_PATH = os.path.join(BASE_DIR, 'dashboard', 'data.json')

if not os.path.exists(EXCEL_PATH):
    raise SystemExit(
        f"Excel bulunamadi: {EXCEL_PATH}\n"
        f"Dosyayi {os.path.join(BASE_DIR, 'data')} altina koyun ya da DMF_EXCEL ile yol verin."
    )

try:
    import openpyxl
except ImportError:
    raise SystemExit("openpyxl kurulu degil.  Kurulum:  pip install openpyxl")

print("Excel yükleniyor...")
wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)

# ── Cycle Times (saniye / parça) ─────────────────────────────────────────────
hat_default_ct = {
    'DMF1': 47, 'DMF2': 49, 'DMF3': 46, 'DMF4': 47,
    'PFW1': 50, 'PFW2': 50, 'PFW3': 50, 'PFW4': 50, 'ITL': 56
}
teorik_max_shift = {h: 480*60/ct for h, ct in hat_default_ct.items()}

ALL_HATS = ['DMF1','DMF2','DMF3','DMF4','PFW1','PFW2','PFW3','PFW4','ITL']

def safe_float(v):
    try: return float(v) if v is not None and isinstance(v, (int,float)) else 0.0
    except: return 0.0

def get_week(ds):
    dt = datetime.strptime(ds, '%Y-%m-%d')
    y, w, _ = dt.isocalendar()
    return f'{y}-H{w:02d}'

# ── DMF Extraction ───────────────────────────────────────────────────────────
print("DMF hatları işleniyor...")

daily_prod  = {}   # {hat: {date: {prod, durus, rework, iskarta}}}
trp_shifts  = {}   # {hat: [{date,vardiya,vt,teorik}]}
vardiya_raw = {}   # {hat: {(date,vardiya): toplam_adet}} — ilk non-zero
ref_prod    = {}   # {hat: {ref: toplam_adet}}
ref_lookup  = {}   # {hat: {(date,vardiya): ref}} — duruş için

DMF_DC = [(7,'Çay/Yemek'),(8,'Temizlik'),(9,'Metod/Proje'),
           (10,'Parça Yok'),(11,'SetUp'),(12,'Tuzak Parça'),
           (13,'Kalite Prob.'),(14,'Arıza')]

for hat in ['DMF1','DMF2','DMF3','DMF4']:
    ws = wb[hat]
    daily_prod[hat]  = {}
    trp_shifts[hat]  = []
    vardiya_raw[hat] = {}
    ref_prod[hat]    = {}
    ref_lookup[hat]  = {}

    for row in ws.iter_rows(min_row=3, values_only=True):
        if not (row[0] and isinstance(row[0], datetime)): continue
        ds = row[0].strftime('%Y-%m-%d')
        v  = str(row[2]) if row[2] and str(row[2]) in ('A','B','C') else None
        ref = str(row[3]).strip() if row[3] else None

        if ds not in daily_prod[hat]:
            daily_prod[hat][ds] = {'prod':0,'durus':defaultdict(float),'rework':0,'iskarta':0}
        d = daily_prod[hat][ds]

        prod = safe_float(row[4])
        if prod > 0:
            d['prod'] += prod
            if ref and ref != 'None' and ' ' not in ref:
                ref_prod[hat][ref] = ref_prod[hat].get(ref,0) + prod

        for ci,cn in DMF_DC:
            if len(row)>ci: d['durus'][cn] += safe_float(row[ci])
        if len(row)>18: d['rework']  += safe_float(row[18])
        if len(row)>20: d['iskarta'] += safe_float(row[20])

        # Referans: son saat satırında (saat 8/16/24) bulunur; boşluk içeren değerler (ör. "Çalışma Yok") atlanır
        if ref and ref != 'None' and v and ' ' not in ref:
            ref_lookup[hat][(ds, v)] = ref

        # Vardiya Toplam: ilk non-zero olan satırı al
        vt = safe_float(row[5])
        if v and vt > 0:
            key = (ds, v)
            if key not in vardiya_raw[hat]:
                vardiya_raw[hat][key] = vt
                trp_shifts[hat].append({'date':ds,'vardiya':v,'vt':vt,'teorik':teorik_max_shift[hat]})

# ── PFW Extraction ───────────────────────────────────────────────────────────
print("PFW hatları işleniyor...")

PFW_DC = [(13,'Çay/Yemek'),(14,'Temizlik'),(15,'Kor.Bakım'),(16,'Metod/Proje'),
           (17,'Sipariş Yok'),(18,'Parça Yok'),(19,'SetUp'),(20,'Kesici Değ.'),
           (21,'Kalite Prob.'),(22,'Arıza')]

for hat in ['PFW1','PFW2','PFW3','PFW4']:
    ws = wb[hat]
    daily_prod[hat]  = {}
    trp_shifts[hat]  = []
    vardiya_raw[hat] = {}
    ref_prod[hat]    = {}

    for row in ws.iter_rows(min_row=3, values_only=True):
        if not (row[0] and isinstance(row[0], datetime)): continue
        ds = row[0].strftime('%Y-%m-%d')
        v  = str(row[3]) if row[3] and str(row[3]) in ('A','B','C') else None
        ref = str(row[4]).strip() if row[4] else None

        if ds not in daily_prod[hat]:
            daily_prod[hat][ds] = {'prod':0,'durus':defaultdict(float),'rework':0,'iskarta':0}
        d = daily_prod[hat][ds]

        if len(row)>10:
            prod = safe_float(row[10])
            if prod > 0:
                d['prod'] += prod
                if ref and ref != 'None':
                    ref_prod[hat][ref] = ref_prod[hat].get(ref,0) + prod

        for ci,cn in PFW_DC:
            if len(row)>ci: d['durus'][cn] += safe_float(row[ci])

        vt = safe_float(row[11]) if len(row)>11 else 0
        if v and vt > 0:
            key = (ds,v)
            if key not in vardiya_raw[hat]:
                vardiya_raw[hat][key] = vt
                trp_shifts[hat].append({'date':ds,'vardiya':v,'vt':vt,'teorik':teorik_max_shift[hat]})

# ── ITL Extraction ───────────────────────────────────────────────────────────
print("ITL hattı işleniyor...")

ITL_DC = [(13,'Çay/Yemek'),(14,'Temizlik'),(15,'Kor.Bakım'),(16,'Metod/Proje'),
           (17,'Sipariş Yok'),(18,'Parça Yok'),(19,'Hat Besleme'),
           (20,'Kalite Prob.'),(21,'Arıza')]

ws = wb['ITL']
daily_prod['ITL']  = {}
trp_shifts['ITL']  = []
vardiya_raw['ITL'] = {}
ref_prod['ITL']    = {}

for row in ws.iter_rows(min_row=4, values_only=True):
    if not (row[0] and isinstance(row[0], datetime)): continue
    ds = row[0].strftime('%Y-%m-%d')
    v  = str(row[2]) if row[2] and str(row[2]) in ('A','B','C') else None
    ref = str(row[3]).strip() if row[3] else None

    if ds not in daily_prod['ITL']:
        daily_prod['ITL'][ds] = {'prod':0,'durus':defaultdict(float),'rework':0,'iskarta':0}
    d = daily_prod['ITL'][ds]

    if len(row)>10:
        prod = safe_float(row[10])
        if prod > 0:
            d['prod'] += prod
            if ref and ref != 'None':
                ref_prod['ITL'][ref] = ref_prod['ITL'].get(ref,0) + prod

    for ci,cn in ITL_DC:
        if len(row)>ci: d['durus'][cn] += safe_float(row[ci])

    vt = safe_float(row[11]) if len(row)>11 else 0
    if v and vt > 0:
        key = (ds,v)
        if key not in vardiya_raw['ITL']:
            vardiya_raw['ITL'][key] = vt
            trp_shifts['ITL'].append({'date':ds,'vardiya':v,'vt':vt,'teorik':teorik_max_shift['ITL']})

# ── SMF-DOOSAN Extraction ────────────────────────────────────────────────────
print("SMF-DOOSAN işleniyor...")

smf_daily = {}  # {date: {jx22, eb2, doosan, total, nok}}
ws = wb['SMF-DOOSAN']
for row in ws.iter_rows(min_row=3, values_only=True):
    if not (row[0] and isinstance(row[0], datetime)): continue
    # Cols: 0=Tarih, 1=JX22-A, 2=JX22-B, 3=JX22-C, 4=NOK, 5=JX22 toplam
    #       6=EB2-A, 7=EB2-B, 8=EB2-C, 9=EB2 toplam, 10=Total
    #       11=Doosan-A, 12=Doosan-B, 13=Doosan-C, 14=Doosan toplam
    jx22 = safe_float(row[5])
    eb2  = safe_float(row[9])
    nok  = safe_float(row[4])
    doos = safe_float(row[14]) if len(row)>14 else 0
    total = jx22 + eb2
    if total > 0:
        ds = row[0].strftime('%Y-%m-%d')
        smf_daily[ds] = {'jx22':jx22, 'eb2':eb2, 'doosan':doos, 'total':total, 'nok':nok}

# ── Duruş Records ────────────────────────────────────────────────────────────
print("Duruş kayıtları işleniyor...")

ws_dur = wb['Duruş Süre&Sebep']
durus_kayitlar = []
for row in ws_dur.iter_rows(min_row=2, values_only=True):
    if not (row[0] and isinstance(row[0], datetime)): continue
    sure = safe_float(row[4])
    if sure <= 0: continue
    hat_d  = str(row[3]) if row[3] else ''
    vrd_d  = str(row[1]) if row[1] else ''
    tar_d  = row[0].strftime('%Y-%m-%d')
    ref_d  = ref_lookup.get(hat_d, {}).get((tar_d, vrd_d), '')
    durus_kayitlar.append({
        'tarih':   tar_d,
        'vardiya': vrd_d,
        'sebep':   str(row[2]) if row[2] else 'Bilinmiyor',
        'hat':     hat_d,
        'ref':     ref_d,
        'sure':    round(sure,1),
        'aciklama':str(row[5])[:90] if row[5] else '',
        'op':      str(row[6]) if row[6] else '',
    })
durus_kayitlar.sort(key=lambda x: x['tarih'], reverse=True)

# ── Weekly & Monthly Aggregates ──────────────────────────────────────────────
print("Özetler hesaplanıyor...")

weekly_prod = {h:{} for h in ALL_HATS}
weekly_trp  = {h:{} for h in ALL_HATS}

for hat in ALL_HATS:
    for ds,d in daily_prod[hat].items():
        if d['prod']==0: continue
        wk = get_week(ds)
        weekly_prod[hat][wk] = weekly_prod[hat].get(wk,0) + d['prod']
    for sh in trp_shifts[hat]:
        wk = get_week(sh['date'])
        if wk not in weekly_trp[hat]: weekly_trp[hat][wk] = []
        weekly_trp[hat][wk].append(sh['vt']/sh['teorik'])

all_weeks = sorted({wk for h in ALL_HATS for wk in weekly_prod[h]})[-20:]

ay_labels = sorted({ds[:7] for h in ALL_HATS for ds,d in daily_prod[h].items() if d['prod']>0})[-8:]

monthly_prod = {}
for hat in ALL_HATS:
    monthly_prod[hat] = [
        int(sum(d['prod'] for ds,d in daily_prod[hat].items() if ds.startswith(ay)))
        for ay in ay_labels
    ]

# ── Vardiya Monthly Averages ─────────────────────────────────────────────────
# Per hat, per month: average shift production for A, B, C
vardiya_monthly = {}   # {hat: {ay: {A:[], B:[], C:[]}}}
for hat in ALL_HATS:
    vardiya_monthly[hat] = {}
    for (ds,v), vt in vardiya_raw[hat].items():
        ay = ds[:7]
        if ay not in vardiya_monthly[hat]:
            vardiya_monthly[hat][ay] = {'A':[],'B':[],'C':[]}
        if v in ('A','B','C'):
            vardiya_monthly[hat][ay][v].append(vt)

vardiya_chart = {}  # {hat: {A:[ort/ay], B:[ort/ay], C:[ort/ay]}}
for hat in ALL_HATS:
    vardiya_chart[hat] = {'A':[],'B':[],'C':[]}
    for ay in ay_labels:
        ayd = vardiya_monthly[hat].get(ay, {'A':[],'B':[],'C':[]})
        for v in ('A','B','C'):
            lst = ayd.get(v,[])
            vardiya_chart[hat][v].append(round(sum(lst)/len(lst),0) if lst else 0)

# ── Reference Distribution ───────────────────────────────────────────────────
# Top 8 refs per hat
ref_top = {}
for hat in ALL_HATS:
    sorted_refs = sorted(ref_prod[hat].items(), key=lambda x: -x[1])[:8]
    ref_top[hat] = {
        'labels': [r[0] for r in sorted_refs],
        'data':   [int(r[1]) for r in sorted_refs],
    }

# TRP per reference: join trp_shifts with ref_lookup
ref_trp_lists = {hat: defaultdict(list) for hat in ['DMF1','DMF2','DMF3','DMF4']}
for hat in ['DMF1','DMF2','DMF3','DMF4']:
    for s in trp_shifts[hat]:
        ref = ref_lookup[hat].get((s['date'], s['vardiya']))
        if ref:
            ref_trp_lists[hat][ref].append(s['vt'] / s['teorik'])

ref_trp_avg = {}
for hat in ['DMF1','DMF2','DMF3','DMF4']:
    entries = [(ref, round(sum(vals)/len(vals)*100, 1), len(vals))
               for ref, vals in ref_trp_lists[hat].items() if len(vals) >= 3]
    entries.sort(key=lambda x: -x[1])
    ref_trp_avg[hat] = {
        'labels': [e[0] for e in entries],
        'trp':    [e[1] for e in entries],
        'shifts': [e[2] for e in entries],
    }

# ── KPI ──────────────────────────────────────────────────────────────────────
kpi = {}
for hat in ALL_HATS:
    all_prod = sum(d['prod'] for d in daily_prod[hat].values())
    son_ay   = sum(d['prod'] for ds,d in daily_prod[hat].items() if ds>='2026-06-01')
    son4 = [t for wk in all_weeks[-4:] for t in weekly_trp[hat].get(wk,[])]
    avg_trp = round(sum(son4)/len(son4),3) if son4 else None
    kpi[hat] = {
        'toplam': int(all_prod), 'son_ay': int(son_ay),
        'rework': int(sum(d['rework'] for d in daily_prod[hat].values())),
        'iskarta': int(sum(d['iskarta'] for d in daily_prod[hat].values())),
        'ort_trp': avg_trp,
    }

# SMF KPI
smf_kpi = {
    'jx22_toplam': int(sum(v['jx22'] for v in smf_daily.values())),
    'eb2_toplam':  int(sum(v['eb2']  for v in smf_daily.values())),
    'toplam':      int(sum(v['total'] for v in smf_daily.values())),
    'nok':         int(sum(v['nok']  for v in smf_daily.values())),
    'aktif_gun':   len(smf_daily),
}

# ── Duruş Pareto ─────────────────────────────────────────────────────────────
durus_cat  = defaultdict(float)
durus_hat2 = defaultdict(float)
durus_aylik= defaultdict(lambda: defaultdict(float))

for r in durus_kayitlar:
    durus_cat[r['sebep']]    += r['sure']
    durus_hat2[r['hat']]     += r['sure']
    durus_aylik[r['tarih'][:7]][r['sebep']] += r['sure']

pareto_cat = sorted(durus_cat.items(), key=lambda x:-x[1])
pareto_hat = sorted(durus_hat2.items(), key=lambda x:-x[1])
aylar      = sorted(durus_aylik.keys())

# ── Weekly series ─────────────────────────────────────────────────────────────
def wk_avg(trp_dict, wk):
    lst = trp_dict.get(wk,[])
    return round(sum(lst)/len(lst)*100,1) if lst else 0

w_prod = {h:[int(weekly_prod[h].get(wk,0)) for wk in all_weeks] for h in ALL_HATS}
w_trp  = {h:[wk_avg(weekly_trp[h],wk) for wk in all_weeks]      for h in ALL_HATS}

# ── SMF daily series ──────────────────────────────────────────────────────────
smf_dates  = sorted(smf_daily.keys())
smf_jx22   = [smf_daily[d]['jx22']  for d in smf_dates]
smf_eb2    = [smf_daily[d]['eb2']   for d in smf_dates]
smf_total  = [smf_daily[d]['total'] for d in smf_dates]

# ── Build JSON ────────────────────────────────────────────────────────────────
print("JSON hazırlanıyor...")

data_json = {
    'meta':   {'guncelleme': datetime.now().strftime('%d.%m.%Y %H:%M'), 'hatlar': ALL_HATS},
    'kpi':    kpi,
    'smf_kpi': smf_kpi,
    'smf_dates': smf_dates,
    'smf_jx22':  smf_jx22,
    'smf_eb2':   smf_eb2,
    'smf_total': smf_total,
    'haftalik_etiket': all_weeks,
    'haftalik_uretim': w_prod,
    'haftalik_trp':    w_trp,
    'aylik_etiket':    ay_labels,
    'aylik_uretim':    monthly_prod,
    'vardiya_chart':   vardiya_chart,
    'ref_top':         ref_top,
    'ref_trp_avg':     ref_trp_avg,
    'pareto_cat_etiket': [x[0] for x in pareto_cat],
    'pareto_cat_deger':  [round(x[1],0) for x in pareto_cat],
    'pareto_hat_etiket': [x[0] for x in pareto_hat],
    'pareto_hat_deger':  [round(x[1],0) for x in pareto_hat],
    'durus_aylar':  aylar,
    'durus_aylik':  {ay:{s:round(v,0) for s,v in sd.items()} for ay,sd in durus_aylik.items()},
    'son_duruslar': durus_kayitlar,
}

# ── Cache ─────────────────────────────────────────────────────────────────────
with open(CACHE_PATH, 'w', encoding='utf-8') as f:
    json.dump(data_json, f, ensure_ascii=False, indent=2)

print(f"\n✓ {CACHE_PATH}")
print(f"  Hatlar:   {', '.join(ALL_HATS)}")
print(f"  Haftalar: {all_weeks[0]} → {all_weeks[-1]}")
print(f"  Duruş kaydı: {len(durus_kayitlar)}")
print("\nDashboard'u üretmek için: python3 generate_dashboard.py")
